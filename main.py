from PySide6.QtWidgets import (QWidget, QVBoxLayout, QPushButton, QFileDialog,
                               QMessageBox, QApplication, QMainWindow, QStatusBar,
                               QLabel, QHBoxLayout)

from classes import TreasureWindow
from functions import (find_workbooks, check_workbook, check_worksheet)
import sys
import json
import os
import pandas as pd
import numpy as np
import openpyxl

DATA_DIRECTORY = 'my_data'
CONDITIONS = f'{DATA_DIRECTORY}/conditions.json'
DAMAGE_TYPES = f'{DATA_DIRECTORY}/damage_types.json'
HIGHLIGHTING = f'{DATA_DIRECTORY}/highlighting.json'
CONFIG_FILE = f'{DATA_DIRECTORY}/config.json'


class StartWindow(QMainWindow):
    def __init__(self, conditions_fp=CONDITIONS, damage_types_fp=DAMAGE_TYPES,
                 highlighting_fp=HIGHLIGHTING, config_file=CONFIG_FILE):
        print(f"main: Starting StartWindow.__init__().")
        super().__init__()
        # Initialize to None any attributes handled by other methods,
        # such as init_ui, load_tables, load_config_files.
        self.conditions_fp = conditions_fp
        self.damage_types_fp = damage_types_fp
        self.highlighting_fp = highlighting_fp
        self.config_fp = config_file
        self.workbook_fps = []
        self.workbook_names = []
        self.conditions = None
        self.damage_types = None
        self.highlighting = None
        self.tables = None
        self.config = None
        self.hbox = None
        self.required_tables = None

        # Create statusbar.
        self.statusbar = QStatusBar()
        self.setStatusBar(self.statusbar)

        # Setting up the attributes for additional windows.
        self.treasure_window = None

        # Starting up the UI.
        self.load_config_files()
        self.load_tables()
        self.init_ui()
        print(f"init: Init process completed.")

    def load_tables(self):
        print(f"load_tables: Starting StartWindow.load_tables().")
        self.tables = {}
        print(f"load_tables: StartWindow.workbook_fps: {self.workbook_fps}.")
        print(f"load_tables: StartWindow.workbook_names: {self.workbook_names}.")
        workbook_checks = find_workbooks(self.workbook_fps)
        if workbook_checks['missing'] is not None:
            missing_txt = ', '.join(str(e) for e in workbook_checks['missing'])
            error_msg = f"Required workbooks; {missing_txt}; could not be found."
            QMessageBox.critical(self, "Fatal Error", error_msg)
            self.exit_app()
        if workbook_checks['corrupt'] is not None:
            corrupt_txt = ', '.join(str(e) for e in workbook_checks['corrupt'])
            error_msg = f"Required workbooks; {corrupt_txt}; could not be opened."
            QMessageBox.critical(self, "Fatal Error", error_msg)
            self.exit_app()

        # Now, we have to tabulate all errors found in the workbooks, such
        # as missing or corrupt worksheets. Extra worksheets will be ignored.
        # The reason is that I often have "Blank" worksheets to make easier on
        # eyes, non-printable "dark theme" worksheets.
        missing_worksheets = {}
        corrupt_worksheets = {}
        print(f"load_tables: missing_worksheets: {missing_worksheets}. "
              f"corrupt_worksheets: {corrupt_worksheets}.")
        print(f"load_tables: StartWindow.config: {self.config}.")
        for idx, wb_name in enumerate(self.workbook_names):
            wb_fp = self.workbook_fps[idx]
            ws_list = self.config['tables']['required tables'][wb_name]
            worksheet_checks = check_workbook(wb_fp, ws_list)
            if worksheet_checks['missing'] is not None:
                missing_worksheets[wb_fp] = worksheet_checks['missing']
            if worksheet_checks['corrupt'] is not None:
                corrupt_worksheets[wb_fp] = worksheet_checks['corrupt']
        if missing_worksheets != {} or corrupt_worksheets != {}:
            missing_txt = ""
            corrupt_txt = ""
            error_msg = ""
            if missing_worksheets != {}:
                for wb_fp in missing_worksheets.keys():
                    ws_list = ' '.join(str(e) for e in missing_worksheets[wb_fp])
                    missing_txt = f"For workbook, {wb_fp}, {ws_list}. {missing_txt}"
                error_msg = f"Missing required worksheets, listed by workbook: " \
                            f"{missing_txt}."

            if corrupt_worksheets != {}:
                for wb_fp in corrupt_worksheets.keys():
                    ws_list = ' '.join(str(e) for e in corrupt_worksheets[wb_fp])
                    corrupt_txt = f"For workbook, {wb_fp}, {ws_list}. {corrupt_txt}"
                error_msg = f"Corrupt required worksheets, listed by workbook: " \
                            f"{corrupt_txt}. {error_msg}"

            QMessageBox.critical(self, "Fatal Error", error_msg)
            self.exit_app()
        # Now, after checking everything thoroughly, we can load the worksheets
        # into tables.
        for idx, wb_name in enumerate(self.workbook_names):
            self.tables[wb_name] = {}
            wb_fp = self.workbook_fps[idx]
            ws_list = self.config['tables']['required tables'][wb_name]
            f = pd.ExcelFile(wb_fp)
            for ws_name in ws_list:
                df = pd.read_excel(wb_fp, sheet_name=ws_name, index_col=0,
                                   na_values=True)
                self.tables[wb_name][ws_name] = df.replace(to_replace=np.nan,
                                                           value=None)
            f.close()
        print(f"load_tables: tables: {self.tables}")

        # Beginning table validation.
        print(f"load_tables: Beginning table validation.")
        bad_worksheets = {}
        errors = 0
        for wb_name in self.tables.keys():
            print(f"load_tables: Starting validation of {wb_name} tables.")
            # The structure of character-related tables differs from the rest.
            # There is an override feature in check_worksheet that handles it.
            stat_override = 'character' in wb_name.lower()
            bad_worksheets[wb_name] = None
            bad_ws_list = []
            for ws_name in self.tables[wb_name].keys():
                if check_worksheet(self.tables[wb_name][ws_name], stat_override):
                    print(f"load_tables: Validated worksheet, {ws_name}.")
                else:
                    bad_ws_list.append(ws_name)
                    errors += 1
                    print(f"load_tables: Worksheet, {ws_name}, has invalid formatting.")

            if len(bad_ws_list) != 0:
                bad_worksheets[wb_name] = bad_ws_list
                print(f"load_tables: bad_worksheets: {bad_worksheets}")
            else:
                print(f"load_tables: Workbook, {wb_name}, has been fully validated.")

        if errors != 0:
            error_msg = ""
            invalid_txt = ""
            for wb_name in bad_worksheets.keys():
                if bad_worksheets[wb_name] is not None:
                    ws_list = ', '.join(str(e) for e in bad_worksheets[wb_name])
                    invalid_txt = f"For workbook, {wb_name}, these worksheets have an "\
                                  f"invalid format: {ws_list}. {invalid_txt}"
            error_msg = f"Invalid formatting of worksheets found, listed by workbook: "\
                        f"{invalid_txt}"
            QMessageBox.critical(self, "Fatal Error", error_msg)
            self.exit_app()
        else:
            print(f"load_tables: All worksheets are valid and ready to use.")
            print(f"load_tables: Completed StartWindow.load_tables().")

    def check_config_file(self, filepath):
        print(f"check_config_file: Starting StartWindow.check_config_file().")
        content = ""
        try:
            with open(filepath, "r") as f:
                content = f.read()
        except (FileNotFoundError, IOError):
            error_msg = f"Configuration file, {filepath}, could not be" \
                        f" read or does not exist."
            QMessageBox.critical(self, "Fatal Error", error_msg)
            self.exit_app()

        if content != "":
            try:
                json_content = json.loads(content)
            except json.decoder.JSONDecodeError:
                error_msg = f"Configuration file, {filepath}, is not a " \
                            f"valid JSON file or may be corrupted."
                QMessageBox.critical(self, "Fatal Error", error_msg)
                self.exit_app()
            else:
                return json_content
        print(f"check_config_file: Completed StartWindow.check_config_file().")

    def load_config_files(self):
        print(f"load_config_files: Starting StartWindow.load_config_files().")
        self.config = self.check_config_file(self.config_fp)
        print(f"load_config_files: config: {self.config}")
        self.conditions = self.check_config_file(self.conditions_fp)
        print(f"load_config_files: conditions: {self.conditions}")
        self.damage_types = self.check_config_file(self.damage_types_fp)
        print(f"load_config_files: damage_types: {self.damage_types}")
        self.highlighting = self.check_config_file(self.highlighting_fp)
        print(f"load_config_files: highlighting: {self.highlighting}")
        # Check config.json to see if it has required tables and workbook dictionary.
        print(f"load_config_files: Checking configuration file.")
        try:
            required_fp_dict = self.config["tables"]['required tables']
        except KeyError:
            error_msg = f"Config file, {self.config_fp} is missing " \
                        f"section for required tables."
            QMessageBox.critical(self, "Fatal Error", error_msg)
            self.exit_app()
        else:
            try:
                dir_fp = DATA_DIRECTORY
                self.workbook_fps = [f"{dir_fp}/{wb_name}" for wb_name in
                                     required_fp_dict.keys()]
                self.workbook_names = [wb_name for wb_name in required_fp_dict.keys()]
            except AttributeError:
                error_msg = f"Config file, {self.config_fp} is missing required" \
                            f"dictionary of workbooks."
                QMessageBox.critical(self, "Fatal Error", error_msg)
                self.exit_app()
        print(f"load_config_files: Configuration file passed checks.")
        print(f"load_config_files: Completed StartWindow.load_config_files().")

    def init_ui(self):
        print(f"init_ui: Starting StartWindow.init_ui().")
        self.setWindowTitle("NPC Generator with Treasures")
        self.setMinimumSize(400, 300)

        treasure_generator_button = QPushButton("Open Treasure Window")
        treasure_generator_button.clicked.connect(self.start_treasure_window)
        exit_button = QPushButton("Exit")
        exit_button.clicked.connect(self.exit_app)
        self.setCentralWidget(QWidget(self))
        self.hbox = QHBoxLayout()
        self.centralWidget().setLayout(self.hbox)
        self.hbox.addWidget(treasure_generator_button)
        self.hbox.addWidget(exit_button)
        print(f"init_ui: Completed StartWindow.init_ui().")

    def start_treasure_window(self):
        print(f"start_treasure_window: Starting StartWindow.start_treasure_window().")
        self.statusbar.showMessage("Opening Treasure Generation Window")
        self.treasure_window = TreasureWindow(self.config, self.conditions,
                                              self.damage_types, self.highlighting,
                                              self.tables)
        self.treasure_window.show()
        print(f"start_treasure_window: Completed StartWindow.start_treasure_window().")

    def exit_app(self):
        sys.exit()


if __name__ == "__main__":
    sys.argv += ['-platform', 'windows:darkmode=2']
    app = QApplication(sys.argv)
    app.setStyle('Fusion')
    window = StartWindow()
    window.show()
    sys.exit(app.exec())
